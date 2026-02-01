import openpyxl
from pathlib import Path
from services.scraper_service import get_product_service
from entities.product import Product
import openpyxl

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "db" / "db.xlsx"

def get_product(id):
    wb = openpyxl.load_workbook(DB_PATH)
    ws = wb.active
    
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if int(row[0]) == int(id):
            prod = Product(
                row[1],
                row[4],
                row[2],
                row[3]
            )
            # id está sendo posto aqui, pois o campo de atributo id só foi criado após a criação do db
            prod.id=row[0]
            wb.close()
            return prod
    wb.close()
    raise ValueError(f"Produto com id {id} não encontrado")

def create_product(name):
    wb = openpyxl.load_workbook(DB_PATH)
    ws = wb.active
    
    prod = get_product_service(name)
    prod.id=ws[f"A{ws.max_row}"].value+1
    ws.append([prod.id,prod.name,prod.lowest_price,prod.seller_name,prod.link])
    wb.save(DB_PATH)
    
    return prod

def update_product(id):
    wb = openpyxl.load_workbook(DB_PATH)
    ws = wb.active
    
    
    for row in ws.iter_rows(min_row=2, values_only=False):
        if int(row[0].value) == int(id):
            prod = get_product_service(row[1].value)
            row[2].value=prod.lowest_price
            row[3].value=prod.seller_name
            row[4].value=prod.link
            # id está sendo posto aqui, pois o campo de atributo id só foi criado após a criação do db
            prod.id=row[0].value
            wb.save(DB_PATH)
            return prod
        
    
    wb.close()
    raise ValueError(f"Produto com id {id} não encontrado")

def delete_product(id):
    wb = openpyxl.load_workbook(DB_PATH)
    ws = wb.active

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True),start=2):
       if int(row[0]) == int(id):
            ws.delete_rows(i)
            wb.save(DB_PATH)
            wb.close()
            return True

    wb.close()
    raise ValueError(f"Produto com id {id} não encontrado")
    



